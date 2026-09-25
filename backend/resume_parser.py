import pypdf
import docx
import re
import os
import traceback
from pathlib import Path

def extract_text_from_pdf(file_path) -> str:
    text = ""
    try:
        with open(str(file_path), "rb") as f:
            reader = pypdf.PdfReader(f)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
    except Exception as e:
        print(f"[Parser ERROR] Failed to parse PDF: {type(e).__name__}: {e}")
    return text.strip()

def extract_text_from_docx(file_path) -> str:
    text_blocks = []
    try:
        doc = docx.Document(str(file_path))
        for para in doc.paragraphs:
            if para.text.strip():
                text_blocks.append(para.text.strip())
                
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text.strip():
                        row_text.append(cell.text.strip())
                if row_text:
                    text_blocks.append(" | ".join(row_text))
                    
    except Exception as e:
        print(f"[Parser ERROR] Failed to parse DOCX: {type(e).__name__}: {e}")
        
    final_text = "\n".join(text_blocks)
    final_text = re.sub(r'\n\s*\n', '\n\n', final_text)
    return final_text.strip()

def extract_text_from_txt(file_path) -> str:
    try:
        with open(str(file_path), "r", encoding="utf-8") as f:
            return f.read().strip()
    except UnicodeDecodeError:
        try:
            with open(str(file_path), "r", encoding="latin-1") as f:
                return f.read().strip()
        except Exception as e:
            print(f"[Parser ERROR] Failed to parse TXT (latin-1 fallback): {type(e).__name__}: {e}")
            return ""
    except Exception as e:
        print(f"[Parser ERROR] Failed to parse TXT: {type(e).__name__}: {e}")
        return ""

def extract_resume_text(file_path) -> str:
    if not file_path:
        return ""
        
    file_path = str(file_path)
    
    print(f"[Parser] Received path: {file_path}")
    print(f"[Parser] Exists: {os.path.exists(file_path)}")
    
    ext = file_path.lower().split(".")[-1]
    print(f"[Parser] Extension: .{ext}")
    
    extracted_text = ""
    if ext == "pdf":
        extracted_text = extract_text_from_pdf(file_path)
    elif ext in ["docx", "doc"]:
        extracted_text = extract_text_from_docx(file_path)
    elif ext in ["txt", "md", "rtf"]:
        extracted_text = extract_text_from_txt(file_path)
        
    print(f"[Parser] Extracted characters: {len(extracted_text)}")
    return extracted_text
